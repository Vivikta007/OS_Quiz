"""Quick checks for the AI question generator and admin access, with fake
Gemini replies (no API key, database or internet needed).

Run:  .\\venv\\Scripts\\python.exe test_question_generator.py
"""

import os

from fastapi.testclient import TestClient

import main
import question_generator as qg
from question_generator import (
    CheckedAnswer,
    CheckedBatch,
    GeneratedBatch,
    GeneratedQuestion,
    generate_questions,
    problems_with,
)


def q(text, options=("A", "B", "C", "D"), correct=0):
    return GeneratedQuestion(
        question=text, options=list(options), correct_index=correct, explanation="Why."
    )


# --- Rule checks -----------------------------------------------------------
assert problems_with(q("Fine?")) == []
assert "has 3 options instead of 4" in problems_with(q("Short?", ("A", "B", "C")))
assert "has duplicate options" in problems_with(q("Dup?", ("A", "a", "C", "D")))
assert "correct answer is out of range" in problems_with(q("Range?", correct=4))
assert "has an empty option" in problems_with(q("Empty?", ("A", " ", "C", "D")))
print("rule checks: ok")


# --- generate_questions with a fake Gemini ---------------------------------
def fake_gemini(writer_batch, checker_batch):
    """Replace the real Gemini call: 1st call = writer, 2nd = checker."""
    replies = [writer_batch, checker_batch]

    def fake_ask(client, instructions, prompt, schema, thinking="medium"):
        reply = replies.pop(0)
        if isinstance(reply, Exception):
            raise reply
        assert isinstance(reply, schema)
        return reply

    qg._ask = fake_ask
    qg._client = lambda: None


fake_gemini(
    GeneratedBatch(questions=[
        q("Checker agrees?", correct=1),
        q("Checker disagrees?", correct=0),
        q("Already exists?"),                   # duplicate of an existing one
        q("Bad options?", ("A", "B", "C")),     # fails the rule checks
    ]),
    CheckedBatch(answers=[
        CheckedAnswer(number=1, answer_index=1),
        CheckedAnswer(number=2, answer_index=2),
    ]),
)
created, skipped = generate_questions("Deadlocks", "Hard", 4, ["Already exists?"])

assert [c["question"] for c in created] == ["Checker agrees?", "Checker disagrees?"]
assert all(c["status"] == "pending" and c["source"] == "ai" for c in created)
assert all(c["topic"] == "Deadlocks" and c["difficulty"] == "Hard" for c in created)
assert created[0]["review_note"] is None
assert 'AI checker chose "C"' in created[1]["review_note"]
assert len(skipped) == 2
assert any("duplicates an existing question" in s for s in skipped)
assert any("has 3 options" in s for s in skipped)
print("generate + check: ok")

# The checker failing must not lose the generated questions
fake_gemini(
    GeneratedBatch(questions=[q("Still saved?")]),
    qg.GenerationError("checker down"),
)
created, _ = generate_questions("Deadlocks", "Easy", 1, [])
assert len(created) == 1 and "Not double-checked" in created[0]["review_note"]
print("checker failure: ok")

# Big requests are split into parts that run in parallel; a failed part
# keeps the others, and the same question from two parts is kept once
def parallel_ask(client, instructions, prompt, schema, thinking="medium"):
    assert thinking == "low"  # Medium difficulty thinks less to be faster
    if schema is CheckedBatch:
        count = prompt.count("\n   0:")
        return CheckedBatch(answers=[CheckedAnswer(number=n, answer_index=0) for n in range(1, count + 1)])
    if "part 3 of 3" in prompt:
        raise qg.GenerationError("part down")
    part = 1 if "part 1 of 3" in prompt else 2
    return GeneratedBatch(questions=[q(f"Part {part} question {n}?") for n in range(4)] + [q("Same in both?")])

qg._ask = parallel_ask
created, skipped = generate_questions("Paging", "Medium", 11, [])
texts = [c["question"] for c in created]
assert len(texts) == 9 and len(set(texts)) == 9, texts
assert any("part down" in s for s in skipped)
assert any("duplicates another new question" in s for s in skipped)

def all_fail(*args, **kwargs):
    raise qg.GenerationError("all down")
qg._ask = all_fail
try:
    generate_questions("Paging", "Hard", 10, [])
    raise AssertionError("expected GenerationError")
except qg.GenerationError as error:
    assert "all down" in str(error)
print("parallel parts: ok")

# From a document: PDFs are attached as a document block, text is inlined,
# and the AI's topics are matched to existing ones (new ones are kept)
calls = []

def document_ask(client, instructions, prompt, schema, thinking="medium"):
    calls.append(prompt)
    if schema is CheckedBatch:
        return CheckedBatch(answers=[])
    return GeneratedBatch(questions=[
        GeneratedQuestion(question="From the doc 1?", topic="cpu  scheduling",
                          options=["A", "B", "C", "D"], correct_index=0, explanation="."),
        GeneratedQuestion(question="From the doc 2?", topic="Disk Scheduling",
                          options=["A", "B", "C", "D"], correct_index=1, explanation="."),
        GeneratedQuestion(question="No topic?", topic=" ",
                          options=["A", "B", "C", "D"], correct_index=1, explanation="."),
    ])

qg._ask = document_ask
pdf = qg.Document(name="notes.pdf", pdf=b"%PDF-1.4 fake")
created, skipped = generate_questions(
    None, "Easy", 3, [], document=pdf, known_topics=["CPU Scheduling", "Paging"]
)
assert [c["topic"] for c in created] == ["CPU Scheduling", "Disk Scheduling"], created
assert any("has no topic" in s for s in skipped)
writer_input = calls[0]
assert writer_input[0]["type"] == "document" and writer_input[0]["mime_type"] == "application/pdf"
assert "CPU Scheduling, Paging" in writer_input[1]["text"]  # existing topics offered
assert isinstance(calls[1], list)  # the checker gets the document too

calls.clear()
notes = qg.Document(name="notes.txt", text="Paging splits memory into frames.")
created, _ = generate_questions("Paging", "Easy", 2, [], document=notes)
assert all(c["topic"] == "Paging" for c in created)  # a fixed topic always wins
assert "Paging splits memory into frames." in calls[0] and "Topic: Paging" in calls[0]
print("from a document: ok")

# A missing API key gives a clear error instead of a crash
os.environ.pop("GEMINI_API_KEY", None)
import importlib
importlib.reload(qg)
qg.load_dotenv = lambda *args, **kwargs: None  # don't read the real backend/.env
try:
    qg._client()
    raise AssertionError("expected GenerationError")
except qg.GenerationError as error:
    assert "GEMINI_API_KEY" in str(error)
print("missing key: ok")


# --- Only admins reach the admin endpoints ---------------------------------
class FakeUser:
    def __init__(self, app_metadata):
        self.app_metadata = app_metadata


client = TestClient(main.app)
for metadata in ({}, {"role": "user"}, None):
    main.app.dependency_overrides[main.get_current_user] = lambda m=metadata: FakeUser(m)
    for method, path, body in (
        ("get", "/admin/questions", None),
        ("post", "/admin/generate", {"topic": "Deadlocks", "difficulty": "Easy"}),
        ("patch", "/admin/questions/1", {"status": "active"}),
    ):
        response = client.request(method, path, json=body)
        assert response.status_code == 403, (metadata, path, response.status_code)
main.app.dependency_overrides.clear()

assert main.is_admin(FakeUser({"role": "admin"}))
assert not main.is_admin(FakeUser(None))
print("admin-only access: ok")


# --- Background generation job (fake Gemini + fake database) ---------------
class FakeQuery:
    def __init__(self, inserted):
        self.inserted = inserted

    def table(self, name):
        return self

    def select(self, columns):
        return self

    def eq(self, column, value):
        return self

    def insert(self, row):
        self.row = row
        return self

    def execute(self):
        if self.row["question"] == "Duplicate?":
            raise main.APIError({"message": "duplicate", "code": "23505"})
        self.inserted.append(self.row)
        return type("Result", (), {"data": [self.row]})()


inserted = []
main.question_bank_client = lambda: FakeQuery(inserted)
main.run_query = lambda query, action: type("Result", (), {"data": []})()
main.generate_questions = lambda *args, **kwargs: (
    [{"question": "New?", "topic": "CPU Scheduling"},
     {"question": "Duplicate?", "topic": "CPU Scheduling"}],
    ["Bad one: has 3 options instead of 4"],
)
main.app.dependency_overrides[main.get_current_user] = lambda: FakeUser({"role": "admin"})

response = client.post(
    "/admin/generate", json={"topic": "  CPU   Scheduling ", "difficulty": "Hard", "count": 2}
)
assert response.status_code == 202, response.text
job_id = response.json()["job_id"]

# TestClient runs background tasks before returning, so the job is finished
job = client.get(f"/admin/generate/{job_id}").json()
assert job["status"] == "done", job
assert job["topic"] == "CPU Scheduling"
assert job["created"] == 1 and [row["question"] for row in inserted] == ["New?"]
assert len(job["skipped"]) == 2 and any("already exists" in s for s in job["skipped"])

def failing_generate(*args, **kwargs):
    raise main.GenerationError("The AI service is unavailable")
main.generate_questions = failing_generate
job_id = client.post("/admin/generate", json={"topic": "X Topic", "difficulty": "Easy"}).json()["job_id"]
job = client.get(f"/admin/generate/{job_id}").json()
assert job["status"] == "failed" and "unavailable" in job["error"]

assert client.get("/admin/generate/unknown-job").status_code == 404
print("background job: ok")


# --- Questions from an uploaded document -----------------------------------
import io

import docx

received = {}

def capture_generate(topic, difficulty, count, existing, guidance, document=None, known_topics=None):
    received.update(topic=topic, document=document, known_topics=known_topics)
    return (
        [{"question": "Doc question 1?", "topic": "Paging"},
         {"question": "Doc question 2?", "topic": "Disk Scheduling"}],
        [],
    )

main.generate_questions = capture_generate
main.run_query = lambda query, action: type("Result", (), {"data": [
    {"question": "Old?", "topic": "Paging", "status": "active"},
    {"question": "Rejected?", "topic": "Old Topic", "status": "rejected"},
]})()
inserted.clear()

def upload(name, content, **form):
    return client.post(
        "/admin/generate-from-document",
        files={"file": (name, content)},
        data={"difficulty": "Medium", "count": "2", **form},
    )

long_text = "Paging divides memory into fixed-size frames. " * 10

response = upload("notes.txt", long_text.encode())
assert response.status_code == 202, response.text
job = client.get(f"/admin/generate/{response.json()['job_id']}").json()
assert job["status"] == "done" and job["created"] == 2 and job["document"] == "notes.txt"
assert job["new_topics"] == ["Disk Scheduling"]  # Paging is already live
assert received["topic"] is None  # no topic sent = let the AI find them
assert received["document"].text.startswith("Paging divides")
assert received["known_topics"] == ["Paging"]  # rejected-only topics aren't offered

word = docx.Document()
word.add_paragraph(long_text)
buffer = io.BytesIO()
word.save(buffer)
response = upload("notes.docx", buffer.getvalue(), topic="  Memory   Management ")
assert response.status_code == 202, response.text
assert received["topic"] == "Memory Management"
assert "Paging divides" in received["document"].text

assert upload("slides.pdf", b"%PDF-1.7 ...").status_code == 202
assert received["document"].pdf.startswith(b"%PDF")

assert upload("virus.exe", b"MZ...").status_code == 415
assert upload("fake.pdf", b"not a pdf").status_code == 400
assert upload("short.txt", b"Too short.").status_code == 400
assert upload("big.txt", b"x" * (10 * 1024 * 1024 + 1)).status_code == 413

main.app.dependency_overrides[main.get_current_user] = lambda: FakeUser({"role": "user"})
assert upload("notes.txt", long_text.encode()).status_code == 403
main.app.dependency_overrides.clear()
print("document upload: ok")

print("\nAll checks passed.")
