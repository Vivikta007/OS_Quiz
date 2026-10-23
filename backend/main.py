import io
import os
import random
import threading
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Literal
from uuid import UUID, uuid4

import httpx
from dotenv import load_dotenv
from fastapi import (
    BackgroundTasks,
    Depends,
    FastAPI,
    File,
    Form,
    Header,
    HTTPException,
    Query,
    Request,
    UploadFile,
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from postgrest.exceptions import APIError
from pydantic import BaseModel, EmailStr, Field
from supabase import Client, ClientOptions, create_client
from supabase_auth.errors import AuthApiError, AuthError, AuthRetryableError

from adaptive import answer_history, pick_adaptive, skill_level
from question_generator import Document, GenerationError, generate_questions

load_dotenv(Path(__file__).resolve().parent / ".env")

app = FastAPI(title="OS Quiz API")

# Browsers block cross-origin requests unless the server allows them.
# Allow the Expo web dev server on localhost or a private LAN address.
app.add_middleware(
    CORSMiddleware,
    allow_origin_regex=(
        r"http://(localhost|127\.0\.0\.1|10\.\d+\.\d+\.\d+|"
        r"192\.168\.\d+\.\d+|172\.(1[6-9]|2\d|3[01])\.\d+\.\d+)(:\d+)?"
    ),
    allow_methods=["*"],
    allow_headers=["*"],
)


# The backend couldn't reach Supabase (timeout, DNS, dropped connection).
# Answer with a clear 503 instead of crashing with a 500 - a crash response
# has no CORS headers, so the app would wrongly report "cannot reach server".
@app.exception_handler(httpx.TransportError)
async def supabase_unreachable(request: Request, error: httpx.TransportError):
    print(f"Cannot reach Supabase: {error!r}", flush=True)
    return JSONResponse(
        status_code=503,
        content={"detail": "The server can't reach the database right now. Try again shortly."},
    )

supabase_url = os.getenv("SUPABASE_URL")
supabase_key = os.getenv("SUPABASE_PUBLISHABLE_KEY")

# The secret key bypasses RLS. It stays in backend/.env only (never in the
# app or on GitHub) and is used ONLY for the questions table, which holds
# the correct answers. User data always goes through the user's own token.
supabase_secret_key = os.getenv("SUPABASE_SECRET_KEY")


# One connection pool shared by every Supabase client. Opening a new secure
# connection to Supabase takes ~0.5s, so reusing open ones makes each request
# much faster. Safe to share: each client sends its own key and user token
# with every request - nothing user-specific is stored on the pool.
# The 8s timeout answers before the app gives up at 10s (REQUEST_TIMEOUT_MS).
supabase_http = httpx.Client(timeout=8.0, follow_redirects=True)


def new_supabase_client(access_token: str | None = None) -> Client:
    """Create a fresh Supabase client for one request.

    A single shared client would remember whoever logged in last, so one
    user's session could leak into another user's request. A fresh client
    per request keeps every user separate. (Only the connection pool,
    supabase_http, is shared - it holds no user tokens.)

    When access_token is given, database queries run *as that user*, so the
    Row Level Security (RLS) policies in Supabase decide what they can see.
    """
    if not supabase_url or not supabase_key:
        raise HTTPException(
            status_code=500,
            detail="Supabase configuration is missing",
        )

    headers = {}
    if access_token:
        headers["Authorization"] = f"Bearer {access_token}"

    return create_client(
        supabase_url,
        supabase_key,
        options=ClientOptions(
            headers=headers,
            auto_refresh_token=False,
            persist_session=False,
            httpx_client=supabase_http,
        ),
    )


def question_bank_client() -> Client:
    """Client for reading the questions table (answers included).

    Only use this for questions - never for user data like quiz_attempts.
    """
    if not supabase_url or not supabase_secret_key:
        print("SUPABASE_SECRET_KEY is missing from backend/.env", flush=True)
        raise HTTPException(
            status_code=500,
            detail="Question bank is not configured on the server",
        )

    return create_client(
        supabase_url,
        supabase_secret_key,
        options=ClientOptions(
            auto_refresh_token=False,
            persist_session=False,
            httpx_client=supabase_http,
        ),
    )


def session_response(response) -> dict:
    """The tokens and user details the app stores after login/refresh."""
    user = response.user
    session = response.session

    return {
        "user_id": str(user.id),
        "email": user.email,
        "full_name": (user.user_metadata or {}).get("full_name"),
        "access_token": session.access_token,
        "refresh_token": session.refresh_token,
        "expires_at": session.expires_at,
    }


# ---------------------------------------------------------------------------
# AUTHENTICATION DEPENDENCIES
# Add `user = Depends(get_current_user)` to any endpoint that needs login.
# ---------------------------------------------------------------------------


def get_access_token(authorization: str | None = Header(default=None)) -> str:
    """Read the token from the 'Authorization: Bearer <token>' header."""
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="Not logged in")

    return authorization.split(" ", 1)[1].strip()


def get_current_user(access_token: str = Depends(get_access_token)):
    """Ask Supabase Auth to verify the token and tell us who the user is.

    Supabase checks the signature and expiry, so the user ID we get back can
    be trusted. Never trust a user_id sent in the request body instead.
    """
    try:
        response = new_supabase_client().auth.get_user(access_token)
    except AuthRetryableError as error:
        print(f"Auth service unavailable: {error}", flush=True)
        raise HTTPException(
            status_code=503,
            detail="Login service unavailable. Try again shortly.",
        ) from error
    except AuthError as error:
        raise HTTPException(
            status_code=401,
            detail="Session expired or invalid",
        ) from error

    if response is None or response.user is None:
        raise HTTPException(status_code=401, detail="Session expired or invalid")

    return response.user


def is_admin(user) -> bool:
    """Admins have "role": "admin" in app_metadata (see migration 005).

    app_metadata can only be changed with the secret key, so users can't
    make themselves admins. (user_metadata, which users CAN edit, is never
    trusted for this.)
    """
    return (user.app_metadata or {}).get("role") == "admin"


def require_admin(user=Depends(get_current_user)):
    """Add `admin=Depends(require_admin)` to endpoints only admins may use."""
    if not is_admin(user):
        raise HTTPException(status_code=403, detail="Admins only")
    return user


# ---------------------------------------------------------------------------
# REQUEST BODIES
# ---------------------------------------------------------------------------


class RegisterRequest(BaseModel):
    full_name: str
    email: EmailStr
    password: str


class LoginRequest(BaseModel):
    email: EmailStr
    password: str


class RefreshRequest(BaseModel):
    refresh_token: str


class ForgotPasswordRequest(BaseModel):
    email: EmailStr


class ResetPasswordRequest(BaseModel):
    email: EmailStr
    # The code from the reset email (6 digits by default; Supabase lets
    # projects choose up to 10)
    code: str = Field(pattern=r"^\d{6,10}$")
    new_password: str


class ProfileUpdateRequest(BaseModel):
    full_name: str = Field(min_length=1, max_length=80)


Difficulty = Literal["Easy", "Medium", "Hard", "Adaptive"]


# What the app sends for each question: which question, the option order
# the user saw, and what they picked. Never the correct answer.
class SubmittedAnswer(BaseModel):
    question_id: int
    options: list[str] = Field(min_length=2, max_length=6)
    selected_answer: str = Field(max_length=500)  # "" = not answered


class QuizAttemptRequest(BaseModel):
    client_attempt_id: UUID
    topics: list[str] = Field(min_length=1, max_length=20)
    difficulty: Difficulty
    answers: list[SubmittedAnswer] = Field(min_length=1, max_length=50)


# ---------------------------------------------------------------------------
# PUBLIC ENDPOINTS
# ---------------------------------------------------------------------------


@app.get("/")
def home():
    return {"message": "OS Quiz Backend is running!"}


@app.get("/health")
def health_check():
    return {"status": "healthy"}


@app.get("/db-check")
def database_check():
    try:
        new_supabase_client().table("profiles").select("id").limit(1).execute()
        return {
            "connected": True,
            "message": "Supabase connection successful",
        }
    except HTTPException:
        raise
    except Exception as error:
        print(f"Database error: {type(error).__name__}: {error}", flush=True)
        raise HTTPException(
            status_code=500,
            detail="Supabase database request failed",
        ) from error


@app.post("/register")
def register_user(data: RegisterRequest):
    if not data.full_name.strip():
        raise HTTPException(
            status_code=400,
            detail="Name cannot be empty",
        )

    if len(data.password) < 8:
        raise HTTPException(
            status_code=400,
            detail="Password must be at least 8 characters",
        )

    try:
        response = new_supabase_client().auth.sign_up({
            "email": str(data.email),
            "password": data.password,
            "options": {
                "data": {
                    "full_name": data.full_name.strip()
                }
            },
        })

        if response.user is None:
            raise HTTPException(
                status_code=400,
                detail="Registration could not be completed",
            )

        return {
            "message": "Registration successful",
            "user_id": str(response.user.id),
            "email": response.user.email,
            "email_confirmation_required": response.session is None,
        }

    except HTTPException:
        raise
    except Exception as error:
        print(
            f"Registration error: {type(error).__name__}: {error}",
            flush=True,
        )
        raise HTTPException(
            status_code=400,
            detail="Registration failed. Check the backend terminal.",
        ) from error


@app.post("/login")
def login_user(data: LoginRequest):
    try:
        response = new_supabase_client().auth.sign_in_with_password({
            "email": str(data.email),
            "password": data.password,
        })
    except AuthRetryableError as error:
        print(f"Login error: {type(error).__name__}: {error}", flush=True)
        raise HTTPException(
            status_code=503,
            detail="Login service unavailable. Try again shortly.",
        ) from error
    except Exception as error:
        print(f"Login error: {type(error).__name__}: {error}", flush=True)
        raise HTTPException(
            status_code=401,
            detail="Invalid email or password, or email not confirmed",
        ) from error

    if response.user is None or response.session is None:
        raise HTTPException(status_code=401, detail="Login failed")

    return {"message": "Login successful", **session_response(response)}


@app.post("/refresh")
def refresh_session(data: RefreshRequest):
    """Swap a refresh token for a new access token.

    Access tokens expire after about an hour. Supabase refresh tokens can
    only be used once, so the app must save the new refresh token we return.
    """
    try:
        response = new_supabase_client().auth.refresh_session(data.refresh_token)
    except AuthRetryableError as error:
        raise HTTPException(
            status_code=503,
            detail="Login service unavailable. Try again shortly.",
        ) from error
    except AuthError as error:
        raise HTTPException(
            status_code=401,
            detail="Session expired. Please log in again.",
        ) from error

    if response.user is None or response.session is None:
        raise HTTPException(
            status_code=401,
            detail="Session expired. Please log in again.",
        )

    return session_response(response)


# ---------------------------------------------------------------------------
# FORGOT PASSWORD
# Step 1: /forgot-password emails the user a one-time code.
# Step 2: /reset-password checks the code and sets the new password.
# A code (not a link) means no deep-link setup is needed in the app, so it
# works the same in Expo Go, on web and in a real build.
# ---------------------------------------------------------------------------


@app.post("/forgot-password")
def forgot_password(data: ForgotPasswordRequest):
    try:
        new_supabase_client().auth.reset_password_for_email(str(data.email))
    except AuthRetryableError as error:
        raise HTTPException(
            status_code=503,
            detail="Login service unavailable. Try again shortly.",
        ) from error
    except AuthApiError as error:
        print(f"Forgot password error: {error.code}: {error}", flush=True)
        if error.status == 429:
            raise HTTPException(
                status_code=429,
                detail="Too many reset emails. Please wait a minute and try again.",
            ) from error
        raise HTTPException(
            status_code=400,
            detail="Could not send the reset email. Please try again.",
        ) from error

    # Same answer whether or not the email has an account, so this endpoint
    # can't be used to find out who is registered.
    return {
        "message": "If an account exists for this email, a reset code has been sent.",
    }


@app.post("/reset-password")
def reset_password(data: ResetPasswordRequest):
    if len(data.new_password) < 8:
        raise HTTPException(
            status_code=400,
            detail="Password must be at least 8 characters",
        )

    # One client for both steps: verifying the code logs this client in as
    # the user, and update_user() then changes *that* user's password.
    client = new_supabase_client()

    try:
        client.auth.verify_otp({
            "email": str(data.email),
            "token": data.code,
            "type": "recovery",
        })
    except AuthRetryableError as error:
        raise HTTPException(
            status_code=503,
            detail="Login service unavailable. Try again shortly.",
        ) from error
    except AuthError as error:
        raise HTTPException(
            status_code=400,
            detail="That code is wrong or has expired. Request a new one.",
        ) from error

    try:
        client.auth.update_user({"password": data.new_password})
    except AuthRetryableError as error:
        raise HTTPException(
            status_code=503,
            detail="Login service unavailable. Try again shortly.",
        ) from error
    except AuthApiError as error:
        print(f"Reset password error: {error.code}: {error}", flush=True)
        if error.code == "same_password":
            detail = "Your new password must be different from the old one."
        elif error.code == "weak_password":
            detail = "That password is too weak. Try a longer one."
        else:
            detail = "Could not change your password. Please try again."
        raise HTTPException(status_code=400, detail=detail) from error

    # Log out every device that used the old password. The password is
    # already changed, so a failure here shouldn't fail the request.
    try:
        client.auth.sign_out({"scope": "global"})
    except Exception as error:
        print(f"Sign out after reset failed: {error}", flush=True)

    return {"message": "Password changed. You can now log in."}


# ---------------------------------------------------------------------------
# PROTECTED ENDPOINTS (login required)
# ---------------------------------------------------------------------------


@app.get("/me")
def get_me(
    user=Depends(get_current_user),
    access_token: str = Depends(get_access_token),
):
    """Return the logged-in user's details and profile."""
    try:
        result = (
            new_supabase_client(access_token)
            .table("profiles")
            .select("full_name, created_at")
            .eq("id", str(user.id))
            .limit(1)
            .execute()
        )
    except Exception as error:
        print(f"Profile error: {type(error).__name__}: {error}", flush=True)
        raise HTTPException(
            status_code=500,
            detail="Could not load your profile",
        ) from error

    profile = result.data[0] if result.data else None

    return {
        "user_id": str(user.id),
        "email": user.email,
        "full_name": profile["full_name"] if profile else None,
        "profile_created_at": profile["created_at"] if profile else None,
        "is_admin": is_admin(user),
    }


# ---------------------------------------------------------------------------
# QUIZ ATTEMPTS (login required)
# Every query filters by the verified user's ID, and RLS in Supabase
# enforces the same rule again, so users only ever see their own attempts.
# ---------------------------------------------------------------------------

ATTEMPT_SUMMARY_COLUMNS = "id, score, total_questions, topics, difficulty, created_at"
ATTEMPT_DETAIL_COLUMNS = f"{ATTEMPT_SUMMARY_COLUMNS}, answers"


def run_query(query, action: str):
    """Run a Supabase query; log the real error, return a safe message."""
    try:
        return query.execute()
    except APIError as error:
        print(f"Database error ({action}): {error}", flush=True)
        raise HTTPException(
            status_code=500,
            detail=f"Could not {action}. Please try again.",
        ) from error


def load_past_answers(access_token: str, user) -> list[dict]:
    """The user's most recent 200 quiz attempts (answers only), oldest first."""
    attempts = run_query(
        new_supabase_client(access_token)
        .table("quiz_attempts")
        .select("answers")
        .eq("user_id", str(user.id))
        .order("created_at", desc=True)
        .limit(200),
        "load your past results",
    ).data
    return list(reversed(attempts))


@app.get("/quiz-questions")
def get_quiz_questions(
    topics: list[str] = Query(min_length=1, max_length=20),
    difficulty: Difficulty = Query(),
    count: int = Query(default=10, ge=1, le=50),
    user=Depends(get_current_user),
    access_token: str = Depends(get_access_token),
):
    """Questions for a new quiz - WITHOUT the correct answers.

    Easy/Medium/Hard: random questions at that difficulty.
    Adaptive: chosen from the user's past results (see adaptive.py).
    """
    query = (
        question_bank_client()
        .table("questions")
        .select("id, question, options, topic, difficulty")
        .eq("status", "active")
        .in_("topic", topics)
    )
    if difficulty != "Adaptive":
        query = query.eq("difficulty", difficulty)

    pool = run_query(query, "load questions").data

    if difficulty == "Adaptive":
        selected = pick_adaptive(pool, load_past_answers(access_token, user), count)
    else:
        selected = random.sample(pool, min(count, len(pool)))

    if not selected:
        raise HTTPException(
            status_code=404,
            detail="No questions found for these topics and difficulty. "
                   "Try another difficulty or add more topics.",
        )

    for question in selected:
        question["options"] = random.sample(
            question["options"], len(question["options"])
        )

    return selected


@app.post("/quiz-attempts", status_code=201)
def create_quiz_attempt(
    data: QuizAttemptRequest,
    user=Depends(get_current_user),
    access_token: str = Depends(get_access_token),
):
    """Mark a finished quiz using the answers in the database, then save it."""
    question_ids = [answer.question_id for answer in data.answers]

    if len(set(question_ids)) != len(question_ids):
        raise HTTPException(status_code=422, detail="Duplicate questions")

    rows = run_query(
        question_bank_client()
        .table("questions")
        .select("id, question, options, correct_index, explanation")
        .in_("id", question_ids),
        "check your answers",
    ).data
    questions_by_id = {row["id"]: row for row in rows}

    # Mark each answer. The correct answer comes from the database,
    # so the app can't change the score.
    #
    # An admin may edit or remove a question while someone is mid-quiz, so
    # the options the player saw can differ from the database by now. That
    # must not throw away their whole quiz: each answer is graded against
    # the answer key as it is NOW, and counts as correct only if it is the
    # current correct option - so sending made-up options can never raise
    # a score.
    graded_answers = []
    for answer in data.answers:
        question = questions_by_id.get(answer.question_id)
        if question is None:
            continue  # deleted from the database since the quiz started

        # The pick must be one of the options this quiz showed
        if answer.selected_answer and answer.selected_answer not in answer.options:
            raise HTTPException(status_code=422, detail="Invalid answer data")

        unchanged = sorted(answer.options) == sorted(question["options"])
        graded_answers.append({
            "question_id": question["id"],
            "question": question["question"],
            # In the order the user saw them - or, if the question was edited
            # during the quiz, the current options (so Review shows the
            # current correct answer)
            "options": answer.options if unchanged else question["options"],
            "selectedAnswer": answer.selected_answer,
            "correctAnswer": question["options"][question["correct_index"]],
            # Shown on the Review screen (older attempts don't have it)
            "explanation": question.get("explanation"),
        })

    if not graded_answers:
        raise HTTPException(
            status_code=422,
            detail="These questions were removed while you were answering. "
                   "Please start a new quiz.",
        )

    score = sum(
        1 for answer in graded_answers
        if answer["selectedAnswer"] == answer["correctAnswer"]
    )

    client = new_supabase_client(access_token)
    row = {
        "user_id": str(user.id),  # from the verified token, never the request
        "client_attempt_id": str(data.client_attempt_id),
        "score": score,
        "total_questions": len(graded_answers),
        "topics": data.topics,
        "difficulty": data.difficulty,
        "answers": graded_answers,
    }

    try:
        result = client.table("quiz_attempts").insert(row).execute()
        saved = result.data[0]
        return {key: saved[key] for key in ATTEMPT_DETAIL_COLUMNS.split(", ")}
    except APIError as error:
        # 23505 = unique violation: this attempt was already saved (the app
        # retried after a network error). Return the saved one instead.
        if getattr(error, "code", None) != "23505":
            print(f"Database error (save quiz attempt): {error}", flush=True)
            raise HTTPException(
                status_code=500,
                detail="Could not save your quiz. Please try again.",
            ) from error

    existing = run_query(
        client.table("quiz_attempts")
        .select(ATTEMPT_DETAIL_COLUMNS)
        .eq("user_id", str(user.id))
        .eq("client_attempt_id", str(data.client_attempt_id))
        .limit(1),
        "save quiz attempt",
    )
    return existing.data[0]


@app.get("/quiz-attempts")
def list_quiz_attempts(
    user=Depends(get_current_user),
    access_token: str = Depends(get_access_token),
):
    """The logged-in user's attempts, newest first (without the answers).

    `topics` lists only the topics the user actually answered questions in.
    The topics picked before the quiz can include ones the quiz never asked
    about (a short quiz doesn't reach every topic), or ones where every
    question was skipped."""
    attempts = run_query(
        new_supabase_client(access_token)
        .table("quiz_attempts")
        .select(ATTEMPT_DETAIL_COLUMNS)
        .eq("user_id", str(user.id))
        .order("created_at", desc=True)
        .limit(100),
        "load your quiz history",
    ).data
    if not attempts:
        return []

    bank = run_query(
        question_bank_client().table("questions").select("id, question, topic"),
        "load your quiz history",
    ).data
    topic_by_id = {row["id"]: row["topic"] for row in bank}
    topic_by_text = {row["question"]: row["topic"] for row in bank}

    for attempt in attempts:
        answered = set()
        for answer in attempt.pop("answers") or []:
            if not answer.get("selectedAnswer"):
                continue  # skipped - not attempted
            # Older attempts saved before question IDs existed: match by text
            topic = topic_by_id.get(answer.get("question_id")) or topic_by_text.get(
                answer.get("question")
            )
            if topic:
                answered.add(topic)
        # Keep the order the user picked them in; topics that were renamed
        # since come last
        attempt["topics"] = [t for t in attempt["topics"] if t in answered] + sorted(
            answered - set(attempt["topics"])
        )
    return attempts


@app.get("/quiz-attempts/{attempt_id}")
def get_quiz_attempt(
    attempt_id: int,
    user=Depends(get_current_user),
    access_token: str = Depends(get_access_token),
):
    """One attempt including its answers, for the Review screen."""
    result = run_query(
        new_supabase_client(access_token)
        .table("quiz_attempts")
        .select(ATTEMPT_DETAIL_COLUMNS)
        .eq("id", attempt_id)
        .eq("user_id", str(user.id))
        .limit(1),
        "load this quiz attempt",
    )

    # Someone else's attempt looks exactly like a missing one (404), so
    # users can't even find out which IDs exist.
    if not result.data:
        raise HTTPException(status_code=404, detail="Quiz attempt not found")

    return result.data[0]


# ---------------------------------------------------------------------------
# PROFILE (login required)
# GET /me returns the profile; PATCH /profile changes it.
# ---------------------------------------------------------------------------


@app.patch("/profile")
def update_profile(
    data: ProfileUpdateRequest,
    user=Depends(get_current_user),
    access_token: str = Depends(get_access_token),
):
    """Change the logged-in user's name. Only full_name can be changed."""
    full_name = data.full_name.strip()

    if not full_name:
        raise HTTPException(status_code=422, detail="Name cannot be empty")

    result = run_query(
        new_supabase_client(access_token)
        .table("profiles")
        .update({"full_name": full_name})
        .eq("id", str(user.id)),
        "update your profile",
    )

    # No row updated = no profile row exists (or RLS blocked it)
    if not result.data:
        raise HTTPException(status_code=404, detail="Profile not found")

    profile = result.data[0]
    return {
        "user_id": str(user.id),
        "email": user.email,
        "full_name": profile["full_name"],
        "profile_created_at": profile["created_at"],
    }


# ---------------------------------------------------------------------------
# PERFORMANCE DASHBOARD (login required)
# ---------------------------------------------------------------------------


def percent(correct: int, total: int) -> int:
    return round(correct / total * 100) if total else 0


def breakdown(stats: dict, order: list[str] | None = None) -> list[dict]:
    """{name: {correct, total}} -> list of rows with an accuracy %."""
    names = order or sorted(stats)
    return [
        {
            "name": name,
            "correct": stats[name]["correct"],
            "total": stats[name]["total"],
            "accuracy": percent(stats[name]["correct"], stats[name]["total"]),
        }
        for name in names
        if name in stats
    ]


@app.get("/performance")
def get_performance(
    user=Depends(get_current_user),
    access_token: str = Depends(get_access_token),
):
    """Summary of the logged-in user's quiz results for the dashboard."""
    attempts = run_query(
        new_supabase_client(access_token)
        .table("quiz_attempts")
        .select(ATTEMPT_DETAIL_COLUMNS)
        .eq("user_id", str(user.id))
        .order("created_at", desc=False)
        .limit(500),
        "load your performance",
    ).data

    if not attempts:
        return {
            "total_attempts": 0,
            "total_questions": 0,
            "total_correct": 0,
            "accuracy": 0,
            "best_score": 0,
            "recent": [],
            "topics": [],
            "difficulties": [],
            "weakest_topic": None,
            "strongest_topic": None,
        }

    # Each saved answer is matched to its question to find its topic and
    # difficulty (Adaptive quizzes mix difficulties, and quizzes can mix
    # topics, so the attempt alone isn't enough). Older attempts saved before
    # question IDs existed are matched by question text instead.
    # Only topic/difficulty are read here - no answers are sent to the app.
    bank = run_query(
        question_bank_client()
        .table("questions")
        .select("id, question, topic, difficulty"),
        "load your performance",
    ).data
    by_id = {row["id"]: row for row in bank}
    by_text = {row["question"]: row for row in bank}

    def new_stats():
        return {"correct": 0, "total": 0}

    topic_stats = defaultdict(new_stats)
    difficulty_stats = defaultdict(new_stats)
    # For the per-topic dropdown in the app: each topic's results by
    # difficulty, and how the user did on that topic in each quiz.
    topic_difficulty_stats = defaultdict(lambda: defaultdict(new_stats))
    topic_quiz_stats = defaultdict(lambda: defaultdict(new_stats))

    for attempt in attempts:
        for answer in attempt["answers"]:
            question = by_id.get(answer.get("question_id")) or by_text.get(
                answer.get("question")
            )
            if question is None:
                continue

            is_correct = answer.get("selectedAnswer") == answer.get("correctAnswer")
            topic = question["topic"]
            for stats in (
                topic_stats[topic],
                difficulty_stats[question["difficulty"]],
                topic_difficulty_stats[topic][question["difficulty"]],
                topic_quiz_stats[topic][attempt["id"]],
            ):
                stats["total"] += 1
                stats["correct"] += int(is_correct)

    total_questions = sum(attempt["total_questions"] for attempt in attempts)
    total_correct = sum(attempt["score"] for attempt in attempts)

    attempts_by_id = {attempt["id"]: attempt for attempt in attempts}
    topics = breakdown(topic_stats)
    for row in topics:
        row["difficulties"] = breakdown(
            topic_difficulty_stats[row["name"]], ["Easy", "Medium", "Hard"]
        )
        # Newest first; the score is for this topic's questions only
        quizzes = sorted(
            topic_quiz_stats[row["name"]].items(),
            key=lambda item: attempts_by_id[item[0]]["created_at"],
            reverse=True,
        )
        row["quizzes"] = [
            {
                "id": attempt_id,
                "created_at": attempts_by_id[attempt_id]["created_at"],
                "difficulty": attempts_by_id[attempt_id]["difficulty"],
                "correct": stats["correct"],
                "total": stats["total"],
                "percent": percent(stats["correct"], stats["total"]),
            }
            for attempt_id, stats in quizzes[:20]
        ]

    # Weakest/strongest only count topics with enough answers to mean something
    ranked = [row for row in topics if row["total"] >= 3]
    weakest = min(ranked, key=lambda row: row["accuracy"], default=None)
    strongest = max(ranked, key=lambda row: row["accuracy"], default=None)

    return {
        "total_attempts": len(attempts),
        "total_questions": total_questions,
        "total_correct": total_correct,
        "accuracy": percent(total_correct, total_questions),
        "best_score": max(
            percent(attempt["score"], attempt["total_questions"])
            for attempt in attempts
        ),
        # The last 10 quizzes, oldest first (for the trend chart)
        "recent": [
            {
                "id": attempt["id"],
                "created_at": attempt["created_at"],
                "score": attempt["score"],
                "total_questions": attempt["total_questions"],
                "difficulty": attempt["difficulty"],
                "percent": percent(attempt["score"], attempt["total_questions"]),
            }
            for attempt in attempts[-10:]
        ],
        "topics": sorted(topics, key=lambda row: -row["accuracy"]),
        "difficulties": breakdown(difficulty_stats, ["Easy", "Medium", "Hard"]),
        "weakest_topic": weakest["name"] if weakest else None,
        "strongest_topic": strongest["name"] if strongest else None,
    }


# ---------------------------------------------------------------------------
# TOPICS (login required) - for the quiz setup screen
# ---------------------------------------------------------------------------


@app.get("/topics")
def list_topics(
    user=Depends(get_current_user),
    access_token: str = Depends(get_access_token),
):
    """Every topic with its question counts and the user's results in it.

    The quiz setup screen uses the counts to grey out empty choices, and the
    accuracy/level to explain what Adaptive mode will focus on.
    """
    rows = run_query(
        question_bank_client()
        .table("questions")
        .select("id, question, topic, difficulty, status"),
        "load topics",
    ).data

    counts = defaultdict(lambda: {"Easy": 0, "Medium": 0, "Hard": 0})
    for row in rows:
        if row["status"] == "active":
            counts[row["topic"]][row["difficulty"]] += 1

    topic_stats, _ = answer_history(load_past_answers(access_token, user), rows)

    result = []
    for topic, level_counts in sorted(counts.items()):
        stats = topic_stats[topic]
        result.append({
            "topic": topic,
            "counts": level_counts,
            "answered": stats["total"],
            "accuracy": percent(stats["correct"], stats["total"])
            if stats["total"]
            else None,
            "level": skill_level(stats["correct"], stats["total"]),
        })
    return result


# ---------------------------------------------------------------------------
# ADMIN: GENERATE AND REVIEW QUESTIONS (admins only)
# Gemini writes new questions as 'pending'. An admin then approves, edits
# or rejects each one. Only 'active' questions ever appear in quizzes.
# These endpoints use the question bank client because they need the
# correct answers - require_admin guards every one of them.
# ---------------------------------------------------------------------------

QuestionStatus = Literal["pending", "active", "rejected"]
ADMIN_QUESTION_COLUMNS = (
    "id, question, options, correct_index, topic, difficulty, status, "
    "source, explanation, review_note, created_at, reviewed_at, version"
)


class GenerateRequest(BaseModel):
    topic: str = Field(min_length=2, max_length=60)
    difficulty: Literal["Easy", "Medium", "Hard"]
    count: int = Field(default=5, ge=1, le=15)
    # Optional steer, e.g. "focus on Round Robin calculations"
    guidance: str = Field(default="", max_length=300)


class QuestionUpdateRequest(BaseModel):
    """Only the fields that are sent get changed.

    expected_version is the question's version when the admin opened it
    (optimistic locking, see migration 006)."""
    expected_version: int = Field(ge=1)
    question: str | None = Field(default=None, min_length=5, max_length=500)
    options: list[str] | None = Field(default=None, min_length=2, max_length=6)
    correct_index: int | None = Field(default=None, ge=0, le=5)
    topic: str | None = Field(default=None, min_length=2, max_length=60)
    difficulty: Literal["Easy", "Medium", "Hard"] | None = None
    explanation: str | None = Field(default=None, max_length=1000)
    status: QuestionStatus | None = None


class QuestionCreateRequest(BaseModel):
    """A question typed in by an admin (not generated)."""
    question: str = Field(min_length=5, max_length=500)
    options: list[str] = Field(min_length=2, max_length=6)
    correct_index: int = Field(ge=0, le=5)
    topic: str = Field(min_length=2, max_length=60)
    difficulty: Literal["Easy", "Medium", "Hard"]
    explanation: str = Field(default="", max_length=1000)
    # Written by an admin, so it can go live straight away
    status: Literal["pending", "active"] = "active"


def is_duplicate_question(error: APIError) -> bool:
    """Postgres 23505 = unique violation (questions.question is unique)."""
    return getattr(error, "code", None) == "23505"


@app.get("/admin/questions")
def admin_list_questions(
    status: QuestionStatus = Query(default="pending"),
    topic: str | None = Query(default=None),
    difficulty: Literal["Easy", "Medium", "Hard"] | None = Query(default=None),
    admin=Depends(require_admin),
):
    """Questions with their answers, for review - plus counts per status."""
    query = (
        question_bank_client()
        .table("questions")
        .select(ADMIN_QUESTION_COLUMNS)
        .eq("status", status)
        .order("id", desc=True)
        .limit(1000)
    )
    if topic:
        query = query.eq("topic", topic)
    if difficulty:
        query = query.eq("difficulty", difficulty)

    questions = run_query(query, "load questions").data

    all_rows = run_query(
        question_bank_client().table("questions").select("status, topic"),
        "count questions",
    ).data
    counts = {"pending": 0, "active": 0, "rejected": 0}
    for row in all_rows:
        counts[row["status"]] += 1

    return {
        "questions": questions,
        "counts": counts,
        "topics": sorted({row["topic"] for row in all_rows}),
        # Topics players can already see. A pending question in any other
        # topic would add a new topic to the app when approved.
        "live_topics": sorted(
            {row["topic"] for row in all_rows if row["status"] == "active"}
        ),
    }


# Generating takes 30-120 seconds (Gemini is slow at busy times), longer
# than the app should wait on one request. So POST /admin/generate starts a
# background job and answers at once; the app polls GET /admin/generate/{id}.
# Jobs live in memory only - a backend restart forgets them, but the
# questions they already saved stay in the database.
generation_jobs: dict[str, dict] = {}
generation_jobs_lock = threading.Lock()
MAX_REMEMBERED_JOBS = 50


def update_job(job_id: str, **fields) -> None:
    with generation_jobs_lock:
        generation_jobs[job_id].update(fields)


def run_generation_job(
    job_id: str,
    topic: str | None,
    difficulty: str,
    count: int,
    guidance: str,
    document: Document | None = None,
) -> None:
    """Runs in the background: generate, check, save as 'pending'.

    topic=None (documents only) lets the AI pick each question's topic."""
    try:
        rows = run_query(
            question_bank_client().table("questions").select("question, topic, status"),
            "load existing questions",
        ).data
        # Every existing question in the topic (any status), so Gemini doesn't
        # repeat one that's live, waiting for review, or already rejected.
        # Without a fixed topic, every question counts, newest first.
        existing = [
            row["question"] for row in reversed(rows)
            if topic is None or row["topic"] == topic
        ]
        # Topics the AI should reuse instead of inventing near-duplicates
        known_topics = sorted({row["topic"] for row in rows if row["status"] != "rejected"})
        live_topics = {row["topic"] for row in rows if row["status"] == "active"}

        generated, skipped = generate_questions(
            topic,
            difficulty,
            count,
            existing,
            guidance,
            document=document,
            known_topics=known_topics,
        )

        # Insert one at a time so a single duplicate doesn't lose the batch.
        created = 0
        new_topics = set()
        for question in generated:
            try:
                question_bank_client().table("questions").insert(question).execute()
                created += 1
                if question["topic"] not in live_topics:
                    new_topics.add(question["topic"])
            except APIError as error:
                if not is_duplicate_question(error):
                    raise
                skipped.append(f"{question['question'][:80]}: already exists")

        update_job(
            job_id,
            status="done",
            created=created,
            skipped=skipped,
            new_topics=sorted(new_topics),
        )
    except GenerationError as error:
        update_job(job_id, status="failed", error=str(error))
    except Exception as error:
        print(f"Generation job failed: {type(error).__name__}: {error}", flush=True)
        update_job(
            job_id,
            status="failed",
            error="Something went wrong while saving the questions. "
                  "Check the backend terminal.",
        )


def start_generation_job(
    background_tasks: BackgroundTasks,
    topic: str | None,
    difficulty: str,
    count: int,
    guidance: str,
    document: Document | None = None,
) -> str:
    job_id = uuid4().hex

    with generation_jobs_lock:
        # Forget the oldest jobs so memory doesn't grow forever
        while len(generation_jobs) >= MAX_REMEMBERED_JOBS:
            generation_jobs.pop(next(iter(generation_jobs)))
        generation_jobs[job_id] = {
            "status": "running",
            "topic": topic,
            "difficulty": difficulty,
            "count": count,
            "document": document.name if document else None,
            "created": 0,
            "skipped": [],
            "new_topics": [],
            "error": None,
        }

    background_tasks.add_task(
        run_generation_job, job_id, topic, difficulty, count, guidance, document
    )
    return job_id


def tidy_topic(topic: str) -> str:
    """Tidy spacing: "  CPU   Sched " -> "CPU Sched"."""
    return " ".join(topic.split())


@app.post("/admin/generate", status_code=202)
def admin_generate_questions(
    data: GenerateRequest,
    background_tasks: BackgroundTasks,
    admin=Depends(require_admin),
):
    """Start generating questions; returns a job id to poll."""
    job_id = start_generation_job(
        background_tasks,
        tidy_topic(data.topic),
        data.difficulty,
        data.count,
        data.guidance.strip(),
    )
    return {"job_id": job_id}


# Questions from an uploaded document (lecture notes, a textbook chapter...).
# PDFs go to Gemini as they are; .txt/.md/.docx are sent as their text.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
DOCUMENT_TYPES = {".pdf", ".txt", ".md", ".docx"}


def read_document(upload: UploadFile) -> Document:
    name = Path(upload.filename or "document").name
    extension = Path(name).suffix.lower()
    if extension not in DOCUMENT_TYPES:
        raise HTTPException(
            status_code=415,
            detail="Upload a PDF, Word (.docx), .txt or .md file.",
        )

    data = upload.file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="The file is bigger than 10 MB.")
    if not data:
        raise HTTPException(status_code=400, detail="The file is empty.")

    if extension == ".pdf":
        if not data.startswith(b"%PDF"):
            raise HTTPException(status_code=400, detail="This file isn't a valid PDF.")
        return Document(name=name, pdf=data)

    if extension == ".docx":
        try:
            import docx  # python-docx; imported here so it's only needed for uploads

            word = docx.Document(io.BytesIO(data))
        except Exception as error:
            raise HTTPException(
                status_code=400, detail="This Word file couldn't be read."
            ) from error
        lines = [paragraph.text for paragraph in word.paragraphs]
        for table in word.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(line for line in lines if line.strip())
    else:
        text = data.decode("utf-8", errors="replace")

    if len(text.strip()) < 200:
        raise HTTPException(
            status_code=400,
            detail="The document has too little text to write questions from.",
        )
    return Document(name=name, text=text)


@app.post("/admin/generate-from-document", status_code=202)
def admin_generate_from_document(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    difficulty: Literal["Easy", "Medium", "Hard"] = Form(...),
    count: int = Form(default=10, ge=1, le=15),
    # Empty = let the AI find the topics in the document (new ones included)
    topic: str = Form(default="", max_length=60),
    guidance: str = Form(default="", max_length=300),
    admin=Depends(require_admin),
):
    """Start writing questions from an uploaded document; returns a job id."""
    document = read_document(file)
    topic = tidy_topic(topic)
    if topic and len(topic) < 2:
        raise HTTPException(status_code=422, detail="The topic is too short.")

    job_id = start_generation_job(
        background_tasks,
        topic or None,
        difficulty,
        count,
        guidance.strip(),
        document,
    )
    return {"job_id": job_id}


@app.get("/admin/generate/{job_id}")
def admin_generation_status(job_id: str, admin=Depends(require_admin)):
    """'running', then 'done' (with created/skipped) or 'failed' (with error)."""
    with generation_jobs_lock:
        job = generation_jobs.get(job_id)
        if job is None:
            raise HTTPException(
                status_code=404,
                detail="This generation job is no longer known (the backend "
                       "restarted). Check the Pending tab for its questions.",
            )
        return dict(job)


@app.post("/admin/questions", status_code=201)
def admin_create_question(
    data: QuestionCreateRequest,
    admin=Depends(require_admin),
):
    """Add one question written by hand."""
    options = [option.strip() for option in data.options]
    if any(not option for option in options):
        raise HTTPException(status_code=400, detail="Options can't be empty")
    if len({option.lower() for option in options}) != len(options):
        raise HTTPException(status_code=400, detail="Options must all be different")
    if not 0 <= data.correct_index < len(options):
        raise HTTPException(status_code=400, detail="Pick which option is correct")

    row = {
        "question": " ".join(data.question.split()),
        "options": options,
        "correct_index": data.correct_index,
        "topic": tidy_topic(data.topic),
        "difficulty": data.difficulty,
        "explanation": data.explanation.strip() or None,
        "status": data.status,
        "source": "manual",
    }
    if data.status == "active":
        row["reviewed_at"] = datetime.now(timezone.utc).isoformat()

    try:
        result = question_bank_client().table("questions").insert(row).execute()
    except APIError as error:
        if is_duplicate_question(error):
            raise HTTPException(
                status_code=409,
                detail="A question with exactly this text already exists.",
            ) from error
        print(f"Database error (create question): {error}", flush=True)
        raise HTTPException(
            status_code=500,
            detail="Could not save the question.",
        ) from error

    return result.data[0]


@app.patch("/admin/questions/{question_id}")
def admin_update_question(
    question_id: int,
    data: QuestionUpdateRequest,
    admin=Depends(require_admin),
):
    """Edit a question and/or change its status (approve = 'active').

    Two admins can work on the same question at once. To stop the second
    save silently overwriting the first, the update only applies if the
    question is still at the version the admin opened (expected_version);
    otherwise it answers 409 and the app reloads the question.
    """
    changes = data.model_dump(exclude_unset=True)
    expected_version = changes.pop("expected_version")
    if not changes:
        raise HTTPException(status_code=400, detail="Nothing to change")

    rows = run_query(
        question_bank_client()
        .table("questions")
        .select(ADMIN_QUESTION_COLUMNS)
        .eq("id", question_id)
        .limit(1),
        "load the question",
    ).data
    if not rows:
        raise HTTPException(status_code=404, detail="Question not found")
    current = rows[0]
    if current["version"] != expected_version:
        raise edited_by_someone_else()

    # Tidy text fields
    for field in ("question", "topic"):
        if field in changes:
            changes[field] = " ".join(changes[field].split())
    if "options" in changes:
        changes["options"] = [option.strip() for option in changes["options"]]
    if "explanation" in changes:
        changes["explanation"] = (changes["explanation"] or "").strip() or None

    # Check the question as it will be AFTER the change
    options = changes.get("options", current["options"])
    correct_index = changes.get("correct_index", current["correct_index"])
    if any(not option for option in options):
        raise HTTPException(status_code=400, detail="Options can't be empty")
    if len({option.lower() for option in options}) != len(options):
        raise HTTPException(status_code=400, detail="Options must all be different")
    if not 0 <= correct_index < len(options):
        raise HTTPException(status_code=400, detail="Pick which option is correct")

    if "status" in changes and changes["status"] != current["status"]:
        changes["reviewed_at"] = datetime.now(timezone.utc).isoformat()
    if changes.get("status") == "active":
        # The admin has checked it, so the AI's warning no longer applies
        changes["review_note"] = None

    try:
        # Compare-and-set in ONE statement:
        #   UPDATE questions SET ... WHERE id = :id AND version = :expected
        # Postgres locks the row while it checks and writes, so if another
        # admin's save lands first (between our read above and this write),
        # their trigger has already bumped the version, nothing matches here,
        # and no row is changed.
        result = (
            question_bank_client()
            .table("questions")
            .update(changes)
            .eq("id", question_id)
            .eq("version", expected_version)
            .execute()
        )
    except APIError as error:
        if is_duplicate_question(error):
            raise HTTPException(
                status_code=409,
                detail="Another question already has exactly this text.",
            ) from error
        print(f"Database error (update question): {error}", flush=True)
        raise HTTPException(
            status_code=500,
            detail="Could not save the question.",
        ) from error

    if not result.data:
        raise edited_by_someone_else()
    return result.data[0]


def edited_by_someone_else() -> HTTPException:
    return HTTPException(
        status_code=409,
        detail="Another admin changed this question while you had it open. "
               "Reload to see their version, then try again.",
    )
